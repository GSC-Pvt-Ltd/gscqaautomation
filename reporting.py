"""Turns collected test results into the two-sheet Excel report.

Reads nothing from pytest directly — it takes plain dicts, so the same writer
works from a JSON file, a re-run, or a future CI step.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK = "1A2224"
MUTED = "5E6B6E"
HEAD_FILL = PatternFill("solid", fgColor="EFF3F2")
PASS_FILL = PatternFill("solid", fgColor="E2F0E7")
FAIL_FILL = PatternFill("solid", fgColor="F8E3E0")
SKIP_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D3DDDC")
BORDER = Border(bottom=THIN)

DETAIL_COLUMNS = [
    ("Case ID", 12), ("Suite", 12), ("Module", 20), ("Title", 46),
    ("Priority", 9), ("Status", 10), ("Duration", 10), ("Error", 60),
    ("Evidence", 34), ("Jira", 12), ("Build", 26),
]


def _style_header(ws, row, columns):
    for i, (name, width) in enumerate(columns, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(bold=True, size=9, color=MUTED)
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width


def _status_fill(status):
    s = (status or "").upper()
    if s == "PASS":
        return PASS_FILL
    if s in ("FAIL", "ERROR"):
        return FAIL_FILL
    return SKIP_FILL


def write_report(results, meta, path):
    """results: list of dicts. meta: run-level dict. path: .xlsx to write."""
    wb = Workbook()

    # ---------- Sheet 1: Summary ----------
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40

    ws["A1"] = "Odoo regression run"
    ws["A1"].font = Font(bold=True, size=14, color=INK)

    total = len(results)
    counts = {}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    passed = counts.get("PASS", 0)
    rate = f"{(passed / total * 100):.1f}%" if total else "n/a"

    rows = [
        ("Environment", meta.get("env", "")),
        ("Target", meta.get("url", "")),
        ("Database", meta.get("db", "")),
        ("Build", meta.get("build", "")),
        ("Started", meta.get("started", "")),
        ("Finished", meta.get("finished", "")),
        ("Duration", f"{meta.get('duration', 0):.1f}s"),
        ("", ""),
        ("Total", total),
        ("Passed", passed),
        ("Failed", counts.get("FAIL", 0) + counts.get("ERROR", 0)),
        ("Skipped", counts.get("SKIP", 0)),
        ("Pass rate", rate),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=bool(k), size=10, color=MUTED)
        ws.cell(row=i, column=2, value=v).font = Font(size=10, color=INK)

    # per-module breakdown
    start = 3 + len(rows) + 1
    ws.cell(row=start, column=1, value="By module").font = Font(bold=True, size=11, color=INK)
    _style_header(ws, start + 1, [("Module", 22), ("Passed", 10), ("Failed", 10)])
    modules = {}
    for r in results:
        m = modules.setdefault(r["module"], {"p": 0, "f": 0})
        if r["status"] == "PASS":
            m["p"] += 1
        elif r["status"] in ("FAIL", "ERROR"):
            m["f"] += 1
    for i, (name, m) in enumerate(sorted(modules.items()), start=start + 2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=m["p"])
        cell = ws.cell(row=i, column=3, value=m["f"])
        if m["f"]:
            cell.fill = FAIL_FILL

    # ---------- Sheet 2: Detail ----------
    ds = wb.create_sheet("Detail")
    _style_header(ds, 1, DETAIL_COLUMNS)
    for i, r in enumerate(results, start=2):
        values = [
            r.get("case_id", ""), r.get("suite", ""), r.get("module", ""),
            r.get("title", ""), r.get("priority", ""), r.get("status", ""),
            round(r.get("duration", 0), 2), r.get("error", ""),
            r.get("evidence", ""), r.get("jira", ""), meta.get("build", ""),
        ]
        for col, val in enumerate(values, start=1):
            c = ds.cell(row=i, column=col, value=val)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(col in (4, 8)))
            if col == 6:
                c.fill = _status_fill(r.get("status"))
                c.font = Font(bold=True, size=10)

    ds.freeze_panes = "A2"
    ds.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_COLUMNS))}{max(1, len(results) + 1)}"

    wb.save(path)
    return path
